from __future__ import annotations
from pathlib import Path

import docx

from app.services.ingestion.types import ExtractedDocument, ExtractedSection, ExtractionError


def extract(path: Path, source_filename: str) -> ExtractedDocument:
    try:
        document = docx.Document(str(path))
    except Exception as exc:  # noqa: BLE001 - python-docx raises several exception types for a bad file
        raise ExtractionError(f"Could not open '{source_filename}' as a Word document.") from exc

    sections: list[ExtractedSection] = []
    current_heading: str | None = None
    buffer: list[str] = []
    section_order = 0

    def flush():
        nonlocal buffer, section_order
        text = "\n\n".join(buffer).strip()
        if text:
            sections.append(
                ExtractedSection(
                    section_order=section_order,
                    page_number=None,  # DOCX has no reliable page concept until rendered
                    section_title=current_heading,
                    text=text,
                )
            )
            section_order += 1
        buffer = []

    for paragraph in document.paragraphs:
        style_name = (paragraph.style.name if paragraph.style else "") or ""
        text = paragraph.text.strip()

        if style_name.startswith("Heading") and text:
            flush()
            current_heading = text
            continue

        if text:
            buffer.append(text)

    flush()

    if not sections:
        raise ExtractionError(f"'{source_filename}' has no readable text content.")

    return ExtractedDocument(source_filename=source_filename, sections=sections, page_count=None)

